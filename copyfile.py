from docx import Document
# Function to transform text
def transform_text(text):
    result = ""
    for ch in text:
        if ch == ".":
            result += ","
        elif ch.islower():
            result += ch.upper()
        elif ch.isupper():
            result += ch.lower()
        else:
            result += ch
    return result
# Load source document
input_file = "data.docx"
output_file = "data1.docx"
doc = Document(input_file)
new_doc = Document()
final_output = ""
# Process each paragraph
for para in doc.paragraphs:
    transformed = transform_text(para.text)
    new_doc.add_paragraph(transformed)
    final_output += transformed + "\n"
# Save new document
new_doc.save(output_file)
# Display output
print("Modified Content:\n")
print(final_output)